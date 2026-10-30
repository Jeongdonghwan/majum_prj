"""키워드 파일 위에서부터 N개 → data/site-pages.json

  python -I scripts/build_pages.py [N]   (기본 5100)

- 키워드 원본: keywords/두피문신_전남광주전북_군제외_키워드.xlsx (군 지역 뺀 6,678개, 순위순)
- wave: 위에서부터 4등분 (post-launch 나눠 열기). 빌드 시 OPEN_WAVE 이하만 공개.
"""
import json, os, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
N = int(sys.argv[1]) if len(sys.argv) > 1 else 5100
ws = load_workbook(os.path.join(ROOT, 'keywords', '두피문신_전남광주전북_군제외_키워드.xlsx'), read_only=True)['키워드']
pages, seen = [], set()
for r in ws.iter_rows(min_row=2, values_only=True):
    no, _, kw, _, term, _, region, level, parent, sido = r[:10]
    slug = kw.replace(' ', '-')
    if slug in seen: continue
    seen.add(slug)
    pages.append(dict(kw=kw, slug=slug, term=term, region=region, level=level, parent=parent, sido=sido))
    if len(pages) == N: break
for i, p in enumerate(pages): p['wave'] = 1 + i * 4 // len(pages)
json.dump({'pages': pages}, open(os.path.join(ROOT, 'data', 'site-pages.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
sys.stdout.reconfigure(encoding='utf-8')
print('pages', len(pages), 'last', pages[-1]['kw'])
