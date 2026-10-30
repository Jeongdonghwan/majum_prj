"""두피문신·탈모 지역 키워드 (전남광주통합특별시 · 전북특별자치도) → keywords/두피문신_전남광주전북_키워드.xlsx

  python -I scripts/gen_keywords.py

키워드: data/keywords.json (사용자 제공 원문 그대로, 추가·선별 없음)
  - 시도·시군·구 × wide 전체
  - 동·읍면 × core 6개
지역: data/regions-kr.json (curtain_prj 에서 복사, 2026-10 전국 행정동 정규화)
"""
import json, os, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
load = lambda n: json.load(open(os.path.join(ROOT, 'data', n), encoding='utf-8'))
KW = load('keywords.json')
WIDE = list(dict.fromkeys(KW['wide']))  # 원문 중복(탈모성지·모발이식·두피문신창업비용)만 제거
CORE = list(dict.fromkeys(KW['core']))
SIDOS = ['전남광주통합특별시', '전북특별자치도']
regions = [r for r in load('regions-kr.json')['regions'] if r['sido'] in SIDOS]

rows, seen = [], set()
for r in regions:  # 시도·시군·구 먼저
    if r['level'] in ('시도', '시군', '구'):
        for w in WIDE:
            kw = f'{r["name"]} {w}'
            if kw not in seen:
                seen.add(kw); rows.append([kw, kw.replace(' ', ''), w, '전체 목록', r['name'], r['level'], r['parent'] or r['name'], r['sido']])
for w in CORE:  # 동·읍면은 핵심 키워드 순서대로
    for r in regions:
        if r['level'] in ('동', '읍면'):
            kw = f'{r["name"]} {w}'
            if kw not in seen:
                seen.add(kw); rows.append([kw, kw.replace(' ', ''), w, '핵심 목록', r['name'], r['level'], r['parent'], r['sido']])

HEAD = PatternFill('solid', fgColor='4A2C5E')
wb = Workbook()


def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color='FFFFFF'), HEAD, Alignment(horizontal='center')
    for d in data:
        ws.append(d)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


ws = wb.active; ws.title = '요약'
lv = Counter((r['sido'], r['level']) for r in regions)
kc = Counter((x[7], x[5]) for x in rows)
summ = [[sd, l, lv[(sd, l)], kc[(sd, l)]] for sd in SIDOS for l in ['시도', '시군', '구', '동', '읍면']]
summ += [['합계', '', len(regions), len(rows)], [], ['전체 목록(시군구용)', len(WIDE), '', ''], ['핵심 목록(동·읍면용)', len(CORE), '', '']]
sheet(ws, ['시도', '지역 단위', '지역 수', '키워드 수'], summ, [22, 20, 10, 10])
sheet(wb.create_sheet('키워드'),
      ['순위', '키워드', '붙여쓰기', '단어', '목록', '지역', '지역 단위', '상위 지역', '시도', '월검색수 PC', '월검색수 모바일', '비고'],
      [[i + 1] + x + [None, None, None] for i, x in enumerate(rows)],
      [8, 36, 32, 26, 10, 18, 9, 10, 20, 12, 14, 14])
sheet(wb.create_sheet('사용 키워드'), ['목록', '키워드'], [['전체', w] for w in WIDE] + [['핵심', w] for w in CORE], [8, 30])

os.makedirs(os.path.join(ROOT, 'keywords'), exist_ok=True)
out = os.path.join(ROOT, 'keywords', '두피문신_전남광주전북_키워드.xlsx')
wb.save(out)
sys.stdout.reconfigure(encoding='utf-8')
for s in summ[:11]:
    print(*s)
print('wide', len(WIDE), 'core', len(CORE), 'total', len(rows), '->', out)
